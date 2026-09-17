from pathlib import Path

from docx import Document

from backend.agents.output_renderer.docx_renderer import render_docx
from backend.schemas.schemas import ComposerResult, OutputType


def sample_composer_result() -> ComposerResult:
    return ComposerResult(
        output_type=OutputType.LITERATURE_SURVEY,
        title="AI Research Assistant",
        content=(
            "## Abstract\nThis is the abstract.\n\n"
            "## Introduction\nThis is the introduction."
        ),
        sections={
            "Abstract": "This is the abstract.",
            "Introduction": "This is the introduction.",
            "References": "1. Smith, John. (2025). AI Research.",
        },
        citations_used=["citation-paper-001"],
    )


def test_render_docx_creates_file(tmp_path: Path):
    output_path = tmp_path / "literature_survey.docx"

    result = render_docx(
        sample_composer_result(),
        output_path,
    )

    assert result == output_path
    assert output_path.exists()
    assert output_path.is_file()


def test_render_docx_contains_title_and_sections(tmp_path: Path):
    output_path = tmp_path / "research_output.docx"

    render_docx(
        sample_composer_result(),
        output_path,
    )

    document = Document(output_path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert "AI Research Assistant" in text
    assert "Abstract" in text
    assert "This is the abstract." in text
    assert "Introduction" in text
    assert "This is the introduction." in text
    assert "References" in text


def test_render_docx_preserves_section_order(tmp_path: Path):
    output_path = tmp_path / "ordered.docx"

    render_docx(
        sample_composer_result(),
        output_path,
    )

    document = Document(output_path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    assert text.index("Abstract") < text.index("Introduction")
    assert text.index("Introduction") < text.index("References")


def test_render_docx_creates_parent_directory(tmp_path: Path):
    output_path = tmp_path / "generated" / "nested" / "output.docx"

    render_docx(
        sample_composer_result(),
        output_path,
    )

    assert output_path.exists()


def test_render_docx_rejects_non_docx_path(tmp_path: Path):
    output_path = tmp_path / "output.pdf"

    try:
        render_docx(
            sample_composer_result(),
            output_path,
        )
        assert False, "Expected ValueError"
    except ValueError as exc:
        assert "Output path must have a .docx extension" in str(exc)