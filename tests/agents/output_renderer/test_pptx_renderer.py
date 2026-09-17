from pathlib import Path

from pptx import Presentation

from backend.agents.output_renderer.pptx_renderer import render_pptx
from backend.schemas.schemas import ComposerResult, OutputType


def sample_composer_result() -> ComposerResult:
    return ComposerResult(
        output_type=OutputType.PPT,
        title="AI Research Assistant",
        content="Presentation content.",
        sections={
            "Problem": "This is the problem statement.",
            "Literature Findings": "These are the literature findings.",
            "Own Approach/Architecture": "This is our architecture.",
            "Own Results": "These are our results.",
            "Comparison": "This is the comparison.",
            "Conclusion": "This is the conclusion.",
        },
        citations_used=["citation-paper-001"],
    )


def test_render_pptx_creates_file(tmp_path: Path):
    output_path = tmp_path / "research_presentation.pptx"

    result = render_pptx(
        sample_composer_result(),
        output_path,
    )

    assert result == output_path
    assert output_path.exists()
    assert output_path.is_file()


def test_render_pptx_contains_title_and_sections(tmp_path: Path):
    output_path = tmp_path / "research_presentation.pptx"

    render_pptx(
        sample_composer_result(),
        output_path,
    )

    presentation = Presentation(str(output_path))

    text = "\n".join(
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if hasattr(shape, "text")
    )

    assert "AI Research Assistant" in text
    assert "Problem" in text
    assert "This is the problem statement." in text
    assert "Literature Findings" in text
    assert "Own Approach/Architecture" in text
    assert "Own Results" in text
    assert "Comparison" in text
    assert "Conclusion" in text


def test_render_pptx_creates_parent_directory(tmp_path: Path):
    output_path = (
        tmp_path
        / "generated"
        / "nested"
        / "presentation.pptx"
    )

    render_pptx(
        sample_composer_result(),
        output_path,
    )

    assert output_path.exists()


def test_render_pptx_preserves_section_order(tmp_path: Path):
    output_path = tmp_path / "ordered.pptx"

    render_pptx(
        sample_composer_result(),
        output_path,
    )

    presentation = Presentation(str(output_path))

    slide_titles = [
        slide.shapes.title.text
        for slide in presentation.slides
        if slide.shapes.title is not None
    ]

    assert slide_titles[0] == "AI Research Assistant"
    assert slide_titles.index("Problem") < slide_titles.index(
        "Literature Findings"
    )
    assert slide_titles.index("Literature Findings") < slide_titles.index(
        "Own Approach/Architecture"
    )
    assert slide_titles.index("Own Approach/Architecture") < slide_titles.index(
        "Own Results"
    )
    assert slide_titles.index("Own Results") < slide_titles.index(
        "Comparison"
    )
    assert slide_titles.index("Comparison") < slide_titles.index(
        "Conclusion"
    )


def test_render_pptx_rejects_non_pptx_path(tmp_path: Path):
    output_path = tmp_path / "presentation.pdf"

    try:
        render_pptx(
            sample_composer_result(),
            output_path,
        )
        assert False, "Expected ValueError"
    except ValueError as exc:
        assert "Output path must have a .pptx extension" in str(exc)