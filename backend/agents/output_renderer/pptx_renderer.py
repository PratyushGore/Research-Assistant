from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

from backend.schemas.schemas import ComposerResult


def render_pptx(
    composer_result: ComposerResult,
    output_path: str | Path,
) -> Path:
    """
    Render a ComposerResult into a PPTX presentation.

    The renderer only formats the content supplied by the Composer.
    It does not generate or modify research content.
    """
    output_path = Path(output_path)

    if output_path.suffix.lower() != ".pptx":
        raise ValueError("Output path must have a .pptx extension.")

    output_path.parent.mkdir(parents=True, exist_ok=True)

    presentation = Presentation()

    # Title slide
    title_slide = presentation.slides.add_slide(
        presentation.slide_layouts[0]
    )

    title_slide.shapes.title.text = composer_result.title

    if len(title_slide.placeholders) > 1:
        subtitle = title_slide.placeholders[1]
        subtitle.text = "Multi-Agent AI Research & Publication Assistant"

    # Content slides
    for section_name, section_content in composer_result.sections.items():
        slide = presentation.slides.add_slide(
            presentation.slide_layouts[1]
        )

        slide.shapes.title.text = section_name

        body = slide.placeholders[1]
        body.text = section_content

        # Keep body text readable
        for paragraph in body.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(20)

    presentation.save(output_path)

    return output_path