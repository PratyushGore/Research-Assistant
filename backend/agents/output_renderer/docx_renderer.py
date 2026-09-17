from pathlib import Path

from docx import Document
from docx.shared import Pt

from backend.schemas.schemas import ComposerResult


def render_docx(
    composer_result: ComposerResult,
    output_path: str | Path,
) -> Path:
    """
    Render a ComposerResult into a DOCX document.

    The renderer only formats the content supplied by the Composer.
    It does not generate or modify research content.
    """
    output_path = Path(output_path)

    if output_path.suffix.lower() != ".docx":
        raise ValueError("Output path must have a .docx extension.")

    output_path.parent.mkdir(parents=True, exist_ok=True)

    document = Document()

    # Document title
    title = document.add_heading(composer_result.title, level=0)

    for run in title.runs:
        run.font.size = Pt(20)

    # Add each Composer section
    for section_name, section_content in composer_result.sections.items():
        document.add_heading(section_name, level=1)

        if section_content:
            paragraph = document.add_paragraph(section_content)
            for run in paragraph.runs:
                run.font.size = Pt(11)

    document.save(output_path)

    return output_path